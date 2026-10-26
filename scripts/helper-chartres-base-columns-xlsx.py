"""Writes analysis/chartres-base-columns.xlsx (tracked in git) from the JSON of
scripts/helper-chartres-base-columns.js (one row per column of the sheet Base).
Usage: python -I scripts/helper-chartres-base-columns-xlsx.py [in.json] [out.xlsx]
Reads only our own JSON. The file holds product names; the user tracks it in analysis/ on purpose.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src = sys.argv[1] if len(sys.argv) > 1 else "downloads/chartres-base-columns.json"
dst = sys.argv[2] if len(sys.argv) > 2 else "analysis/chartres-base-columns.xlsx"
data = json.load(open(src, encoding="utf-8"))

HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")
FILLS = {
    "from the file": None,
    "checked on the data": PatternFill("solid", fgColor="D5E8D4"),
    "from the macro code": PatternFill("solid", fgColor="D5E8D4"),
    "inferred, confirm with the site": PatternFill("solid", fgColor="FFE6CC"),
}


def sheet(wb, title, headers, widths, rows, certainty_col):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill = FILLS.get(row[certainty_col].value)
        if fill:
            row[certainty_col].fill = fill
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


wb = Workbook()
wb.remove(wb.active)

readme = wb.create_sheet("Read me")
lines = [
    "Chartres, sheet Base of the OEE entry workbooks: one row per column.",
    "",
    f"Sheet 'Base columns (reference)': the reference template, {data['reference']} ({len(data['refRows'])} columns), in the order of the file.",
    "Sheet 'Other lines (extra columns)': columns of the other 2026 files whose name is not in the reference "
    f"({len(data['extraRows'])} rows). Column letters are those of the first line listed.",
    "",
    "Columns: French = the name in row 1; Entry code = row 2 (ENR1, ENR2, ...); English = translation (machine parts: word by word, check them); "
    "Example = the value of the newest day that has a non-zero value, with that date in brackets; "
    "Formula = the formula of row 3 as stored in the file; Feeds = the report table that would use it; "
    "In lines = in how many of the 18 workbooks of 2026 the exact name exists.",
    "",
    "Written by = the part of the entry form (BSAISIE, VBA read but not run) that writes the column: a form field, a formula the form inserts, or a value it computes. "
    "A column with nothing there is a spare 'Libre' column that the form never writes.",
    "",
    "How sure: 'from the file' = the name or the sheet Category says it; 'from the macro code' = read in the VBA; "
    "'checked on the data' = a rule tested on the files; 'inferred, confirm with the site' = worked out from formulas and numbers.",
    "",
    "The machine part names (Dysf / Panne / Réglages) are different in every line: the English there is a word-by-word translation.",
    "Contains product names from the files and is tracked in git (analysis/): do not push the repo without checking.",
]
for line in lines:
    readme.append([line])
readme.column_dimensions["A"].width = 160
for row in readme.iter_rows():
    row[0].alignment = Alignment(wrap_text=True, vertical="top")

ref_rows = [
    [
        r["index"], r["letter"], r["french"], r["code"], r["en"], r["type"], r["example"],
        r["meaning"], r["formula"], r["writtenBy"], r["feeds"], r["certainty"], f"{r['inLines']} of {r['of']}",
    ]
    for r in data["refRows"]
]
sheet(
    wb, "Base columns (reference)",
    ["#", "Letter", "French (as in the file)", "Entry code", "English", "Type", "Example (newest day with a value)",
     "Meaning", "Formula in the file (row 3)", "Written by (macro code)", "Feeds", "How sure", "In lines"],
    [5, 7, 38, 9, 38, 14, 28, 70, 30, 34, 30, 22, 9],
    ref_rows, 11,
)

extra_rows = [
    [
        r["letter"], r["french"], r["code"], r["en"], r["type"], r["example"], r["meaning"],
        r["formula"], r["writtenBy"], r["feeds"], r["certainty"], r["inLines"], r["lineList"],
    ]
    for r in data["extraRows"]
]
sheet(
    wb, "Other lines (extra columns)",
    ["Letter (first line)", "French (as in the file)", "Entry code", "English", "Type", "Example (first line listed)",
     "Meaning", "Formula in the file (row 3)", "Written by (macro code)", "Feeds", "How sure", "Number of lines", "Lines"],
    [9, 42, 9, 42, 14, 28, 70, 30, 34, 30, 22, 9, 40],
    extra_rows, 10,
)

wb.save(dst)
print(f"wrote {dst}: {len(ref_rows)} reference rows, {len(extra_rows)} extra rows")
